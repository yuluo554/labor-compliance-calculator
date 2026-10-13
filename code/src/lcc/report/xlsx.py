"""XLSX 渲染（可选依赖 ``openpyxl``）——M5 实现（FR-6.1）。

**分层说明**：本模块是**可选依赖边界**——唯一允许 import ``openpyxl`` 的报告侧
文件（守卫式惰性导入）。同层的姊妹模块 :mod:`lcc.report` 只做 Markdown + 章节结构，
保持 stdlib-only（零依赖导出通路）。核心层（engine/params/gate/law/model/pipeline）
仍零第三方依赖（D-09）。

内容与 Markdown **同源**：两者都消费 :func:`lcc.report.report_sections` 的章节结构，
故 FR-6.1 的「两格式内容一致」由单一渲染源保证，而非人工比对。
"""

from __future__ import annotations

import os
import re
import zipfile
from datetime import datetime

from ..errors import OptionalDependencyError

MILESTONE = "M5"

# 固定文档时间戳 → 同一输入导出可复现（docProps 不写入"生成时刻"）
FIXED_DOC_TIMESTAMP = datetime(2026, 1, 1, 0, 0, 0)

# ZIP 容器确定性参数（沿用 D-16「zip 型产物统一固定时间戳 + create_system=0」）
_FIXED_ZIP_DATETIME = (1980, 1, 1, 0, 0, 0)
_ZIP_CREATE_SYSTEM = 0

_CORE_MODIFIED_RE = re.compile(
    rb"(<dcterms:modified[^>]*>)[^<]*(</dcterms:modified>)"
)
_CORE_CREATED_RE = re.compile(
    rb"(<dcterms:created[^>]*>)[^<]*(</dcterms:created>)"
)

_PARAM_FALLBACK = [("（无参数来源登记）", "")]
_WAGE_FALLBACK = [("（不适用）", "")]


def require_openpyxl():
    # type: () -> object
    """守卫式惰性导入 openpyxl；缺失 → :class:`OptionalDependencyError`（exit 1）。"""
    try:
        import openpyxl
        from openpyxl.styles import Alignment, Font
    except ImportError as exc:
        raise OptionalDependencyError(
            "XLSX 导出需要 openpyxl（可选依赖），请先安装：",
            ["python -m pip install -e \".[report]\"",
             "（原因：{}）".format(exc),
             "未安装时请改用 Markdown 导出（无需任何第三方依赖）。"],
        )
    return openpyxl, Alignment, Font


def _canonicalize_archive(path):
    # type: (str) -> None
    """把 openpyxl 写出的 xlsx 重打包为**逐字节可复现**的 zip。

    为什么需要这一步：``openpyxl`` 的 ``save_workbook()`` 会在写盘瞬间无条件覆盖
    ``docProps/core.xml`` 的 ``dcterms:modified`` 为当前 UTC 时刻
    （实测 openpyxl 3.1.5 ``writer/excel.py``），且 zip 条目时间戳取当前时间——
    于是同一输入两次导出**字节不同**，与 D-16 的确定性落盘纪律冲突。

    处理：读出全部条目 → 归一化 core.xml 的两个时间戳 → 以固定时间戳、
    ``create_system=0``、固定压缩方式重写（不依赖 ``ZIP_STORED`` 与否，
    只要参数固定即可复现；沿用 D-16 的约定取值）。
    """
    with zipfile.ZipFile(path, "r") as source:
        entries = [(item.filename, source.read(item.filename)) for item in source.infolist()]

    fixed_iso = FIXED_DOC_TIMESTAMP.strftime("%Y-%m-%dT%H:%M:%SZ").encode("ascii")
    canonical = []
    for name, data in entries:
        if name == "docProps/core.xml":
            data = _CORE_MODIFIED_RE.sub(rb"\g<1>" + fixed_iso + rb"\g<2>", data)
            data = _CORE_CREATED_RE.sub(rb"\g<1>" + fixed_iso + rb"\g<2>", data)
        canonical.append((name, data))

    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as target:
        for name, data in canonical:
            info = zipfile.ZipInfo(name, date_time=_FIXED_ZIP_DATETIME)
            info.compress_type = zipfile.ZIP_DEFLATED
            info.create_system = _ZIP_CREATE_SYSTEM
            info.external_attr = 0o600 << 16
            target.writestr(info, data)


def render_xlsx(result, path, sections=None):
    # type: (object, str, object) -> str
    """渲染 XLSX 测算说明；返回写入的路径。

    :param sections: 可传入 :func:`lcc.report.report_sections` 的结果以复用；
        为 None 时自行计算（保证两格式内容一致）。
    """
    openpyxl, Alignment, Font = require_openpyxl()

    if sections is None:
        from . import report_sections

        sections = report_sections(result)

    workbook = openpyxl.Workbook()
    bold = Font(bold=True)
    wrap = Alignment(wrap_text=True, vertical="top")

    def write_pairs(sheet, pairs, widths=(28, 72)):
        # type: (object, list, tuple) -> None
        for label, value in pairs:
            sheet.append([label, value])
        if sheet.max_row >= 1:
            for row in sheet.iter_rows(min_row=1, max_row=max(sheet.max_row, 1), max_col=2):
                row[0].font = bold
                row[1].alignment = wrap
        sheet.column_dimensions["A"].width = widths[0]
        sheet.column_dimensions["B"].width = widths[1]

    summary = workbook.active
    summary.title = "基本信息"
    write_pairs(summary, sections["meta"])

    write_pairs(workbook.create_sheet("输入事实"),
                sections["inputs"] or _WAGE_FALLBACK, widths=(28, 72))

    write_pairs(workbook.create_sheet("参数来源"),
                sections["params"] or _PARAM_FALLBACK, widths=(36, 96))

    results = workbook.create_sheet("计算结果")
    results.append(["模块", "项目", "内容"])
    for cell in results[1]:
        cell.font = bold
    for module in sections["modules"]:
        if module["fields"]:
            for label, value in module["fields"]:
                results.append([module["label"], label, value])
        else:
            results.append([module["label"], "（无字段）", ""])
    results.column_dimensions["A"].width = 26
    results.column_dimensions["B"].width = 34
    results.column_dimensions["C"].width = 72

    steps = workbook.create_sheet("计算过程")
    steps.append(["模块", "序号", "过程留痕"])
    for cell in steps[1]:
        cell.font = bold
    for module in sections["modules"]:
        for index, step in enumerate(module["steps"], 1):
            steps.append([module["label"], index, step])
    steps.column_dimensions["A"].width = 26
    steps.column_dimensions["B"].width = 8
    steps.column_dimensions["C"].width = 96

    details = [item for item in sections["modules"] if item["detail"]]
    if details:
        sheet = workbook.create_sheet("明细")
        sheet.append(["表格", "模块"] + list(details[0]["detail"]["headers"]))
        for cell in sheet[1]:
            cell.font = bold
        for module in details:
            detail = module["detail"]
            for row in detail["rows"]:
                sheet.append([detail["title"], module["label"]] + list(row))

    basis = workbook.create_sheet("适用依据")
    basis.append(["引用", "标题", "要点"])
    for cell in basis[1]:
        cell.font = bold
    if sections["basis"]:
        for clause in sections["basis"]:
            basis.append([clause["citation"], clause["title"], "\n".join(clause["points"])])
    else:
        basis.append(["（本情形无适用补偿条款）", "", ""])
    basis.column_dimensions["A"].width = 40
    basis.column_dimensions["B"].width = 40
    basis.column_dimensions["C"].width = 72

    notice = workbook.create_sheet("说明与免责声明")
    for item in sections["notices"]:
        notice.append([item])
    if sections["notices"]:
        notice.append([])
    notice.append([sections["disclaimer"]])
    notice.column_dimensions["A"].width = 100

    workbook.properties.created = FIXED_DOC_TIMESTAMP
    workbook.properties.modified = FIXED_DOC_TIMESTAMP

    directory = os.path.dirname(os.path.abspath(path))
    if directory and not os.path.isdir(directory):
        os.makedirs(directory)
    workbook.save(path)
    _canonicalize_archive(path)
    return path


__all__ = [
    "MILESTONE",
    "FIXED_DOC_TIMESTAMP",
    "require_openpyxl",
    "render_xlsx",
]
